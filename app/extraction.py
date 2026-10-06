import hashlib
import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ExtractionError(Exception):
    pass


def extract_text(content, extension):
    try:
        if extension == "txt":
            text = content.decode("utf-8-sig")
            return [(f"líneas {start + 1}-{min(start + 80, len(text.splitlines()))}", "\n".join(text.splitlines()[start:start + 80]))
                    for start in range(0, len(text.splitlines()), 80)]
        if extension == "pdf":
            reader = PdfReader(io.BytesIO(content), strict=True)
            return [(f"página {i}", page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
        if extension == "docx":
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                paragraphs.extend(" | ".join(cell.text.strip() for row in table.rows for cell in row.cells).splitlines())
            return [("documento", "\n".join(paragraphs))]
    except Exception as exc:
        raise ExtractionError("No se pudo leer el archivo. Comprueba que no está dañado o protegido.") from exc
    raise ExtractionError("Formato de archivo no admitido.")


def clean_text(text):
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_pages(pages, chunk_size=1000, overlap=150):
    if chunk_size < 100 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Configuración de fragmentación inválida")
    chunks = []
    step = chunk_size - overlap
    for locator, raw in pages:
        text = clean_text(raw)
        for start in range(0, len(text), step):
            piece = text[start:start + chunk_size].strip()
            if piece:
                chunks.append((locator if len(text) <= chunk_size else f"{locator}, caracteres {start + 1}-{start + len(piece)}", piece,
                               hashlib.sha256(piece.encode("utf-8")).hexdigest()))
    return chunks
